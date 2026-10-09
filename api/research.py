"""Blueprint: research-mode dataset uploads and manifests."""

from __future__ import annotations

import csv
import errno
import io
import json
import logging
import mimetypes
import os
import re
import tempfile
import time
import uuid
import zipfile
from dataclasses import asdict
from datetime import datetime
from functools import wraps
from pathlib import Path
from shutil import disk_usage

from flask import Blueprint, abort, jsonify, request, send_file
from werkzeug.utils import secure_filename

from agent.r_analysis_templates import enriched_analysis_templates
from agent.analysis.catalog import AnalysisCatalog
from agent.analysis.memory import DatasetRoleMemory
from agent.analysis.preparation import AnalysisPreparer
from agent.analysis.readiness import (
    build_analysis_readiness_plan,
    build_template_compatibility_guidance,
    has_format_specific_preparation_route,
)
from agent.analysis.run_control import research_run_controller
from agent.analysis.workflows import default_workflow_registry, validate_workflow_run_spec
from agent.jobs import JobCanceled, JobLimitExceeded
from data.analysis_plan_store import AnalysisPlanStore
from data.research_dataset import ResearchDatasetStore
from data.scientific_dataset import infer_scientific_asset_specs
from data.research_results import (
    SCHEMA_VERSION as RESULT_SCHEMA_VERSION,
    ResearchResultStore,
)
from data.research_workspace import default_research_workspace
from infrastructure.paths import data_path
from infrastructure.public_projection import public_projection, public_r_result
from data.reproducibility import without_replay_snapshots
from r_runtime.config import (
    RRuntimeConfig,
    get_product_mode,
    is_research_mode,
)
from r_runtime.dependency_gate import (
    check_dependency_ticket,
    dependency_ticket,
    inspect_dependency_gate,
)
from r_runtime.environment import HostREnvironment, REnvironmentError
from r_runtime.factory import get_r_backend
from r_runtime.inspection import inspection_summary
from r_runtime.package_policy import dataset_loader_packages
from r_runtime.execution import dataset_inspection_request, run_r_request
from r_runtime.protocol import DatasetAssetRef, DatasetRef
from r_runtime.resource_policy import compute_local_resource_policy
from .local_access import session_for_request


bp = Blueprint("research", __name__)
log = logging.getLogger(__name__)


def _exclusive_research_mutation(view):
    """Serialize result/dataset mutations against an active analysis turn."""

    @wraps(view)
    def wrapped(sid: str, *args, **kwargs):
        sess = session_for_request(sid)
        if not sess.begin_session_mutation():
            return jsonify({"error": "当前会话正在分析或变更状态，请稍后重试。"}), 409
        try:
            return view(sid, *args, **kwargs)
        finally:
            sess.end_session_mutation()

    return wrapped


@bp.errorhandler(PermissionError)
def research_file_permission_error(_error: PermissionError):
    return jsonify({"error": "本地研究数据无法访问"}), 403


_TEMPLATE_RUNTIME_CACHE_TTL_SECONDS = 120
_BROWSER_PREVIEW_MAX_ROWS = 100
_BROWSER_PREVIEW_MAX_COLUMNS = 50
_BROWSER_PREVIEW_MAX_FIELD_BYTES = 2 * 1024
_BROWSER_PREVIEW_MAX_RESPONSE_BYTES = 512 * 1024
_TEMPLATE_PREPARE_MAX_REQUEST_BYTES = 64 * 1024
_BROWSER_UPLOAD_CHUNK_BYTES = 1024 * 1024
_BROWSER_UPLOAD_FREE_RESERVE_BYTES = 64 * 1024 * 1024
_template_runtime_cache: dict[str, object] = {
    "expires_at": 0.0,
    "packages": [],
    "ok": False,
    "error": "",
    "r_version": "",
}


def invalidate_template_runtime_inventory() -> None:
    """Force the next template projection to use the newly selected R library."""

    _template_runtime_cache.update(
        {"expires_at": 0.0, "packages": [], "ok": False, "error": "", "r_version": ""}
    )


def _positive_int_env(name: str, default: int) -> int:
    try:
        return max(1, int(os.environ.get(name, str(default))))
    except ValueError:
        return default


def _research_upload_limit_bytes() -> int:
    """Return zero for an unlimited loopback-only compatibility upload."""
    configured = os.environ.get("BAA_RESEARCH_MAX_FILE_BYTES")
    if configured is None:
        return 0
    try:
        return max(0, int(configured or str(64 * 1024 * 1024)))
    except ValueError:
        return 64 * 1024 * 1024


def _browser_upload_reserve_bytes() -> int:
    try:
        return max(
            0,
            int(
                os.environ.get(
                    "BAA_RESEARCH_UPLOAD_DISK_RESERVE_BYTES",
                    str(_BROWSER_UPLOAD_FREE_RESERVE_BYTES),
                )
            ),
        )
    except ValueError:
        return _BROWSER_UPLOAD_FREE_RESERVE_BYTES


def _ensure_browser_upload_capacity(root: Path, required_bytes: int) -> None:
    """Keep an explicit free-space reserve on the final dataset filesystem."""
    try:
        free_bytes = int(disk_usage(root).free)
    except OSError as exc:
        raise ValueError("无法确认本地数据存储的剩余空间") from exc
    needed = max(0, int(required_bytes)) + _browser_upload_reserve_bytes()
    if free_bytes < needed:
        raise ValueError("本地数据存储空间不足，无法安全接收该上传文件")


def _stream_browser_upload(upload, destination: Path, *, expected_bytes: int, limit: int) -> int:
    """Stream one upload to final-filesystem staging with bounded free-space checks."""
    _ensure_browser_upload_capacity(destination.parent, expected_bytes)
    written = 0
    try:
        with destination.open("xb") as handle:
            while True:
                chunk = upload.stream.read(_BROWSER_UPLOAD_CHUNK_BYTES)
                if not chunk:
                    break
                written += len(chunk)
                if limit and written > limit:
                    raise ValueError("单文件超过大小限制")
                _ensure_browser_upload_capacity(destination.parent, len(chunk))
                handle.write(chunk)
            handle.flush()
            os.fsync(handle.fileno())
    except ValueError:
        destination.unlink(missing_ok=True)
        raise
    except OSError as exc:
        destination.unlink(missing_ok=True)
        if exc.errno in {errno.ENOSPC, getattr(errno, "EDQUOT", errno.ENOSPC)}:
            raise ValueError("本地数据存储空间不足，上传未写入") from exc
        raise ValueError("无法写入本地数据存储") from exc
    return written


def _store() -> ResearchDatasetStore:
    return ResearchDatasetStore(data_path("uploads", "research"))


def _llm_readiness() -> dict:
    from .state import config_manager

    configs = config_manager.list_configs()
    enabled = [
        provider
        for provider, config in configs.items()
        if config.get("enabled") and config.get("has_api_key")
    ]
    default_provider = config_manager.get_default_provider()
    default_config = configs.get(default_provider) or {}
    defaults = config_manager.DEFAULT_CONFIGS["deepseek"]
    configured = bool(default_config.get("has_api_key"))
    return {
        "configured": configured,
        "bundled": False,
        "default_provider": default_provider,
        "model": default_config.get("model") or defaults["model"],
        "display_name": default_config.get("name") or defaults["display_name"],
        "base_url": defaults["base_url"],
        "transport": "https",
        "context_window": default_config.get("context_window") or defaults["context_window"],
        "context_label": defaults["context_label"],
        "max_output_tokens": default_config.get("max_output_tokens") or defaults["max_output_tokens"],
        "max_output_label": defaults["max_output_label"],
        "enable_thinking": default_config.get("enable_thinking", defaults["enable_thinking"]),
        "reasoning_effort": default_config.get("reasoning_effort"),
        "enabled_count": len(enabled),
        "providers": enabled,
    }


def _r_runtime_readiness() -> dict:
    from r_runtime.isolation import isolation_availability

    cfg = RRuntimeConfig.from_env()
    isolation = isolation_availability()
    environment = HostREnvironment(cfg.rscript_bin or None)
    try:
        verified = environment.prepare()
        available = bool(isolation["available"])
        version = verified.r_version
        runtime_id = verified.runtime_id
        package_count = verified.package_count
        package_library = str(verified.install_library)
        library_mode = verified.library_mode
        error = str(isolation["error"])
    except REnvironmentError as exc:
        available = False
        version = ""
        runtime_id = ""
        package_count = 0
        package_library = cfg.package_library
        library_mode = ""
        error = str(exc)

    return {
        "backend": cfg.backend,
        "isolation": isolation,
        "label": "用户本机 R",
        "detail": error or cfg.rscript_bin or "未找到 Rscript",
        "available": available,
        "rscript_bin": cfg.rscript_bin,
        "package_library": package_library,
        "library_mode": library_mode,
        "package_count": package_count,
        "runtime_id": runtime_id,
        "timeout_seconds": cfg.timeout_seconds,
        "worker_count": cfg.worker_count,
        "native_thread_count": cfg.native_thread_count,
        "memory_budget_bytes": cfg.memory_budget_bytes,
        "in_memory_threshold_bytes": cfg.in_memory_threshold_bytes,
        "r_version": version,
        "version_text": f"R {version}" if version else "",
        "error": error,
    }


def _active_job_count(sid: str) -> int:
    try:
        from .state import session_manager

        sess = session_manager.get(sid)
        if sess is None:
            return 0
        return len(sess.job_runner.list_jobs(active_only=True, limit=20, top_level_only=True))
    except Exception:
        return 0


def _capability(id_: str, label: str, status: str, detail: str = "") -> dict:
    return {
        "id": id_,
        "label": label,
        "status": status,
        "detail": detail,
    }


def _dataset_ref(dataset) -> DatasetRef:
    return DatasetRef(
        dataset_id=dataset.dataset_id,
        path=dataset.stored_path,
        format=dataset.format,
        name=dataset.name,
        assets=tuple(
            DatasetAssetRef(
                path=asset.path,
                format=asset.format,
                role=asset.role,
                relative_mount_name=asset.relative_mount_name,
                kind=asset.kind,
                name=asset.name,
                size_bytes=asset.size_bytes,
                device=asset.device,
                inode=asset.inode,
                object_type=asset.object_type,
                mtime_ns=asset.mtime_ns,
                ctime_ns=asset.ctime_ns,
                fingerprint_sha256=asset.fingerprint_sha256,
            )
            for asset in getattr(dataset, "assets", ())
        ),
    )


def _session_datasets(sid: str, _sess) -> list:
    current = default_research_workspace().current(sid)
    return [current] if current is not None else []


def _result_store(sid: str) -> ResearchResultStore:
    return default_research_workspace().results(sid)


def _allowed_artifact_roots(sid: str, sess) -> list[Path]:
    # Results are written under the app-owned output root.  A Scientific
    # Dataset Bundle may reference a user's external directory; treating that
    # directory as an artifact root would turn the preview/download endpoint
    # into a browser for unrelated neighboring files.
    return [data_path("outputs", "research", sid)]


def _resolve_artifact_path(sid: str, raw_path: str, sess) -> Path | None:
    if not raw_path or "\x00" in raw_path:
        return None
    try:
        candidate = Path(raw_path).expanduser().resolve(strict=True)
    except (OSError, RuntimeError):
        return None
    if not candidate.is_file():
        return None

    for root in _allowed_artifact_roots(sid, sess):
        try:
            candidate.relative_to(root.resolve(strict=False))
            return candidate
        except ValueError:
            continue
    return None


def _table_delimiter(path: Path) -> str:
    return "\t" if path.suffix.lower() == ".tsv" else ","


def _is_table_artifact(path: Path) -> bool:
    return path.suffix.lower() in {".csv", ".tsv"}


def _live_template_runtime_inventory() -> dict[str, object]:
    now = time.time()
    if now < float(_template_runtime_cache.get("expires_at") or 0):
        return dict(_template_runtime_cache)
    try:
        payload = get_r_backend().list_packages(timeout_seconds=15)
    except Exception as exc:  # pragma: no cover - defensive runtime boundary
        payload = {"ok": False, "error": str(exc), "packages": []}
    records = payload.get("packages") if isinstance(payload, dict) else []
    package_names = sorted(
        {
            str(record.get("package") or record.get("Package") or "").strip()
            for record in records
            if isinstance(record, dict) and str(record.get("package") or record.get("Package") or "").strip()
        }
    )
    next_value: dict[str, object] = {
        "expires_at": now + _TEMPLATE_RUNTIME_CACHE_TTL_SECONDS,
        "packages": package_names,
        "ok": bool(payload.get("ok")) and bool(package_names),
        "error": "" if payload.get("ok") else str(payload.get("error") or ""),
        "r_version": str(payload.get("r_version") or ""),
    }
    _template_runtime_cache.update(next_value)
    return dict(next_value)


@bp.get("/api/session/<sid>/research/readiness")
def research_readiness(sid: str):
    sess = session_for_request(sid)
    r_runtime = _r_runtime_readiness()
    llm = _llm_readiness()
    datasets = _session_datasets(sid, sess)
    active_jobs = _active_job_count(sid)
    r_status = "ready" if r_runtime["available"] else "needs_setup"
    ai_status = "ready" if llm["configured"] else "needs_setup"

    capabilities = [
        _capability("upload", "上传数据", "ready", "仅保留用户文件上传"),
        _capability("inspect", "理解数据", r_status, r_runtime["label"]),
        _capability("r_execute", "R 执行", r_status, r_runtime["detail"]),
        _capability("artifacts", "结果产物", "ready", "表格、图形、文本输出"),
        _capability("job_runner", "任务队列", "ready", f"{active_jobs} 个运行中任务"),
        _capability(
            "ai",
            "AI 模型",
            ai_status,
            llm["model"] if llm["configured"] else "请配置 DeepSeek API Key",
        ),
        _capability("connectors", "外部数据接入", "disabled", "研究模式已关闭"),
    ]

    return jsonify(
        {
            "ok": True,
            "mode": get_product_mode(),
            "research_mode": is_research_mode(),
            "r_runtime": r_runtime,
            "llm": llm,
            "session": {
                "datasets": len(datasets),
                "active_jobs": active_jobs,
            },
            "capabilities": capabilities,
            "base_modules": {
                "session_manager": True,
                "job_runner": True,
                "tool_registry": True,
                "artifact_api": True,
                "chat_sse": True,
                "llm_config": True,
                "research_upload_store": True,
                "external_connectors_enabled": False,
            },
        }
    )


@bp.get("/api/session/<sid>/research/templates")
def research_templates(sid: str):
    session_for_request(sid)
    runtime = _live_template_runtime_inventory()
    return jsonify(
        enriched_analysis_templates(
            runtime_packages=runtime.get("packages") if isinstance(runtime.get("packages"), list) else [],
            runtime_ok=bool(runtime.get("ok")),
            runtime_error=str(runtime.get("error") or ""),
            runtime_r_version=str(runtime.get("r_version") or ""),
        )
    )


def _template_preparation_profile(sid: str, dataset) -> dict:
    summary = (
        dataset.profile.get("summary")
        if isinstance(dataset.profile, dict) and isinstance(dataset.profile.get("summary"), dict)
        else {}
    )
    record = default_research_workspace().current_record(sid) or {}
    assets = [asset.to_dict() for asset in getattr(dataset, "assets", ())]
    scientific = summary.get("scientific") if isinstance(summary.get("scientific"), dict) else {}
    return {
        "dataset_id": dataset.dataset_id,
        "format": dataset.format,
        "assets": assets,
        "asset_formats": list(
            dict.fromkeys(str(asset.get("format") or "") for asset in assets if asset.get("format"))
        ),
        "object_kind": str(summary.get("object_kind") or "data.frame"),
        "modality": str(summary.get("modality") or scientific.get("modality") or "tabular"),
        "representation": str(summary.get("representation") or ""),
        "processing_stage": str(summary.get("processing_stage") or ""),
        "scientific": scientific,
        "workspace_generation": str(record.get("workspace_generation") or ""),
        "profile_revision": str(summary.get("profile_revision") or ""),
        "rows": summary.get("rows"),
        "columns": summary.get("variables") if isinstance(summary.get("variables"), list) else [],
        "local_semantics": summary.get("local_semantics")
        if isinstance(summary.get("local_semantics"), dict)
        else {},
    }


def _template_workflow_preflight(template_id: str, profile: dict, plan=None):
    workflow_spec = default_workflow_registry().compile(
        plan if plan is not None else template_id,
        dataset_profile=profile,
    )
    return workflow_spec, validate_workflow_run_spec(workflow_spec, profile)


def _compatible_tabular_template_suggestions(profile: dict, *, exclude: str) -> list[dict]:
    scientific = profile.get("scientific")
    scientific = scientific if isinstance(scientific, dict) else {}
    modality = str(profile.get("modality") or scientific.get("modality") or "").casefold()
    representation = str(
        profile.get("representation") or scientific.get("representation") or ""
    ).casefold()
    if modality not in {"", "tabular", "table"} and not any(
        marker in representation for marker in ("data.frame", "data_frame", "dataframe", "tabular")
    ):
        return []
    raw_variables = profile.get("variables")
    if not isinstance(raw_variables, list):
        raw_variables = profile.get("columns")
    variables = raw_variables if isinstance(raw_variables, list) else []
    typed_variables = [item for item in variables if isinstance(item, dict)]
    semantic_types = [str(item.get("semantic_type") or "").casefold() for item in typed_variables]
    numeric_count = semantic_types.count("numeric")
    categorical_count = semantic_types.count("categorical")
    if not categorical_count:
        for item in typed_variables:
            if str(item.get("semantic_type") or "").casefold() != "numeric":
                continue
            try:
                unique_count = int(float(item.get("unique_count")))
            except (TypeError, ValueError):
                continue
            if 2 <= unique_count <= 20:
                categorical_count += 1
    candidates: list[tuple[str, str]] = [
        ("data_quality", "先核对变量类型、缺失、异常和后续可分析性"),
        ("distribution_visual", "探索当前表格中变量分布、异常值和科研图形"),
    ]
    if categorical_count and numeric_count:
        candidates.extend(
            (
                ("publication_table", "按真实分组生成描述性 Table 1"),
                ("group_tests", "比较当前表格中的连续或分类结局"),
            )
        )
    if numeric_count >= 2:
        candidates.extend(
            (
                ("correlation_agreement", "分析当前数值变量的相关性与一致性"),
                ("regression_glm", "根据研究问题指定结局和暴露变量建模"),
            )
        )
    catalog = AnalysisCatalog()
    suggestions: list[dict] = []
    for template_id, reason in candidates:
        if template_id == exclude:
            continue
        try:
            template = catalog.get(template_id)
        except KeyError:
            continue
        suggestions.append({"id": template.id, "title": template.title, "reason": reason})
    return suggestions[:6]


@bp.post("/api/session/<sid>/research/templates/<template_id>/prepare")
@_exclusive_research_mutation
def prepare_research_template(sid: str, template_id: str):
    session_for_request(sid)
    request.max_content_length = _TEMPLATE_PREPARE_MAX_REQUEST_BYTES
    if request.content_length and request.content_length > _TEMPLATE_PREPARE_MAX_REQUEST_BYTES:
        return jsonify({"error": "模板准备请求过大"}), 413
    payload = request.get_json(silent=True) or {}
    if not isinstance(payload, dict):
        return jsonify({"error": "请求格式无效"}), 400
    raw_inputs = payload.get("inputs", {})
    if not isinstance(raw_inputs, dict):
        return jsonify({"error": "模板输入格式无效"}), 400
    dataset_id = str(payload.get("dataset_id") or "").strip()
    dataset = default_research_workspace().current(sid)
    if dataset is None:
        return jsonify({"error": "请先上传数据"}), 404
    if dataset_id and dataset_id != dataset.dataset_id:
        return jsonify({"error": "准备请求与当前数据集不一致，请重新选择模板。"}), 409
    summary = (
        dataset.profile.get("summary")
        if isinstance(dataset.profile, dict) and isinstance(dataset.profile.get("summary"), dict)
        else {}
    )
    if not isinstance(summary.get("variables"), list):
        return jsonify({"error": "请等待当前数据集完成自动理解后再使用模板。"}), 400

    analysis_context = (
        dataset.profile.get("analysis_context")
        if isinstance(dataset.profile.get("analysis_context"), dict)
        else {}
    )
    memory = DatasetRoleMemory.from_dict(analysis_context.get("role_memory"))
    try:
        runtime = _live_template_runtime_inventory()
        live_packages = (
            runtime.get("packages")
            if bool(runtime.get("ok")) and isinstance(runtime.get("packages"), list)
            else None
        )
        prepared = AnalysisPreparer(
            AnalysisCatalog(),
            memory,
            runtime_packages=live_packages,
        ).prepare(
            template_id,
            _template_preparation_profile(sid, dataset),
            raw_inputs,
        )
    except KeyError:
        return jsonify({"error": "分析模板不存在"}), 404
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400

    dataset.profile["analysis_context"] = {**analysis_context, "role_memory": memory.to_dict()}
    _store().save(dataset)
    plan_payload = None
    dependency_gate = None
    preparation_plan = None
    compatibility_guidance = None
    preparation_profile = _template_preparation_profile(sid, dataset)
    workflow_spec, preflight = _template_workflow_preflight(
        template_id,
        preparation_profile,
        prepared.plan,
    )
    question_ids = {item.input_id for item in prepared.questions}
    question_resolved_issue_codes = (
        {"unknown_omics_axes"} if "__matrix_orientation__" in question_ids else set()
    )
    structural_issues = tuple(
        issue
        for issue in preflight.issues
        if issue.code != "unresolved_critical_inputs"
        and issue.code not in question_resolved_issue_codes
    )
    preparation_issues = ()
    if prepared.plan is not None:
        preparation_issues = preflight.issues
    elif prepared.status == "needs_input":
        # Format, object, stage and asset blockers must be handled before the
        # user is asked to map variables that are not yet analysis-ready.
        preparation_issues = structural_issues
    elif prepared.status == "not_applicable" and has_format_specific_preparation_route(
        preparation_profile
    ):
        # Only known source formats receive a conversion route. An unrelated
        # ordinary table cannot be presented as convertible into RNA-seq,
        # imaging or another scientifically different modality.
        preparation_issues = structural_issues
    if preparation_issues:
        preparation_plan = build_analysis_readiness_plan(
            workflow_spec,
            preparation_profile,
            preparation_issues,
        )
    elif prepared.plan is not None:
        AnalysisPlanStore().save(sid, prepared.plan)
        plan_payload = asdict(prepared.plan)
        dependency_gate = inspect_dependency_gate(
            packages=prepared.plan.packages,
            intent=" ".join(
                (
                    prepared.plan.intent,
                    prepared.plan.method_family,
                    *prepared.plan.prerequisites,
                )
            ),
            mode="exact" if AnalysisCatalog().get(template_id).package_policy == "required" else "candidate",
            session_id=sid,
            plan_id=prepared.plan.plan_id,
            template_id=prepared.plan.template_id,
        )
    status_labels = {
        "ready": "可直接使用",
        "needs_input": "需确认关键项",
        "not_applicable": "当前数据不适用",
        "needs_preparation": "需先准备数据",
    }
    response = asdict(prepared)
    response_status = "needs_preparation" if preparation_plan is not None else prepared.status
    response["status"] = response_status
    response["status_label"] = status_labels[response_status]
    if preparation_plan is not None:
        response["reason"] = str(preparation_plan.get("headline") or prepared.reason)
        response["preparation_plan"] = preparation_plan
    elif response_status == "not_applicable":
        compatibility_guidance = build_template_compatibility_guidance(
            workflow_spec,
            preparation_profile,
            prepared.reason,
            alternative_templates=_compatible_tabular_template_suggestions(
                preparation_profile,
                exclude=template_id,
            ),
        )
        response["compatibility_guidance"] = compatibility_guidance
    response["plan"] = plan_payload
    response["dependency_gate"] = dependency_gate
    return jsonify({"ok": True, **response})


@bp.post("/api/session/<sid>/research/dependencies/<ticket_id>/check")
def check_research_dependencies(sid: str, ticket_id: str):
    session_for_request(sid)
    try:
        report = check_dependency_ticket(sid, ticket_id)
    except ValueError as exc:
        return jsonify({"error": str(exc), "code": "dependency_ticket_invalid"}), 404
    return jsonify({"ok": True, "dependency_gate": report})


@bp.post("/api/session/<sid>/research/dependencies/<ticket_id>/install")
def install_research_dependencies(sid: str, ticket_id: str):
    sess = session_for_request(sid)
    try:
        ticket = dependency_ticket(ticket_id, sid)
        report = check_dependency_ticket(sid, ticket_id)
    except ValueError as exc:
        return jsonify({"error": str(exc), "code": "dependency_ticket_invalid"}), 404
    packages = tuple(str(item) for item in report.get("install_packages") or [] if str(item))
    if not packages:
        return jsonify(
            {
                "error": "当前依赖需要按提示配置或修复，不能作为普通 R 包自动安装。",
                "code": "dependency_not_installable",
                "dependency_gate": report,
            }
        ), 409
    if not sess.begin_session_mutation():
        return jsonify({"error": "当前会话正在分析或变更状态，请稍后重试。"}), 409

    def install(ctx):
        from agent.analysis.run_control import ResearchRunCancelled

        research_run_controller.begin(ctx.job_id)
        try:
            ctx.check_canceled()
            ctx.set_progress(5, "正在检查本机 R 与包库")
            environment = HostREnvironment.locate()
            verified = environment.prepare()
            ctx.check_canceled()
            ctx.set_progress(10, "正在等待安装资源：" + "、".join(packages))
            environment.ensure_packages(
                packages, verified=verified, conversation_job_id=ctx.job_id,
                on_output=lambda message: ctx.set_progress(20, message),
            )
            ctx.check_canceled()
            ctx.set_progress(90, "安装进程结束，正在重新验证依赖")
            invalidate_template_runtime_inventory()
            refreshed = check_dependency_ticket(sid, ticket.ticket_id)
            ctx.check_canceled()
            return {"ok": True, "dependency_gate": refreshed}
        except ResearchRunCancelled as exc:
            raise JobCanceled(ctx.job_id) from exc
        finally:
            invalidate_template_runtime_inventory()
            research_run_controller.finish(ctx.job_id)

    try:
        job_id = sess.job_runner.create(
            install, "dependency_install", label="安装 R 包：" + "、".join(packages),
            on_done=sess.end_session_mutation,
        )
    except Exception:
        sess.end_session_mutation()
        raise
    return jsonify({"ok": True, "job_id": job_id, "dependency_gate": report}), 202


@bp.get("/api/session/<sid>/research/result-state")
def get_research_result_state(sid: str):
    sess = session_for_request(sid)
    result_store = _result_store(sid)
    normalized, exists, saved_at = result_store.load()
    if not exists:
        if result_store.generation:
            return jsonify({"ok": True, "exists": False})
        path = result_store.path
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            payload = {}
        state = payload.get("state") if isinstance(payload, dict) else None
        if isinstance(state, dict):
            return jsonify(
                {
                    "ok": True,
                    "exists": True,
                    "saved_at": str(payload.get("saved_at") or ""),
                    "state": without_replay_snapshots(state),
                }
            )
        return jsonify({"ok": True, "exists": False})
    return jsonify(
        {
            "ok": True,
            "exists": True,
            "saved_at": saved_at,
            "schema_version": RESULT_SCHEMA_VERSION,
            "state": without_replay_snapshots(normalized),
        }
    )


@bp.post("/api/session/<sid>/research/result-state")
@_exclusive_research_mutation
def save_research_result_state(sid: str):
    sess = session_for_request(sid)
    if request.content_length and request.content_length > 8 * 1024 * 1024:
        return jsonify({"error": "结果状态过大，无法保存"}), 413
    payload = request.get_json(silent=True) or {}
    if not isinstance(payload, dict):
        return jsonify({"error": "结果状态格式无效"}), 400
    state = payload.get("state") if isinstance(payload.get("state"), dict) else payload
    if (
        state.get("schema_version") == RESULT_SCHEMA_VERSION
        or isinstance(state.get("analysis_turns"), list)
        or isinstance(state.get("paper_collection"), dict)
    ):
        try:
            normalized = _result_store(sid).save(state)
        except ValueError as exc:
            return jsonify({"error": str(exc)}), 409
        return jsonify(
            {
                "ok": True,
                "saved_at": normalized["updated_at"],
                "schema_version": RESULT_SCHEMA_VERSION,
                "state": without_replay_snapshots(normalized),
            }
        )
    path = _result_store(sid).path
    path.parent.mkdir(parents=True, exist_ok=True)
    record = {
        "saved_at": datetime.now().isoformat(timespec="seconds"),
        "state": without_replay_snapshots(state),
    }
    path.write_text(json.dumps(record, ensure_ascii=False), encoding="utf-8")
    return jsonify({"ok": True, "saved_at": record["saved_at"]})


@bp.delete("/api/session/<sid>/research/result-state")
@_exclusive_research_mutation
def clear_research_result_state(sid: str):
    sess = session_for_request(sid)
    _result_store(sid).delete()
    return jsonify({"ok": True})


@bp.post("/api/session/<sid>/research/result-state/reset")
@_exclusive_research_mutation
def reset_research_result_state(sid: str):
    sess = session_for_request(sid)
    result_store = _result_store(sid)
    previous_state, _exists, _saved_at = result_store.load()
    active_dataset_id = str(getattr(sess, "research_dataset_id", "") or "") or str(
        (previous_state.get("dataset") or {}).get("dataset_id") or ""
    )
    state = result_store.reset()
    if active_dataset_id:
        try:
            dataset = _store().get(sid, active_dataset_id)
        except ValueError:
            dataset = None
        if dataset is not None:
            sess.research_dataset_id = dataset.dataset_id
            state = result_store.set_dataset(dataset.to_dict(), dataset.profile.get("summary") or {})
    return jsonify(
        {
            "ok": True,
            "exists": True,
            "saved_at": state["updated_at"],
            "schema_version": RESULT_SCHEMA_VERSION,
            "state": without_replay_snapshots(state),
        }
    )


@bp.patch("/api/session/<sid>/research/artifacts/<artifact_id>/core")
@_exclusive_research_mutation
def patch_research_artifact_core(sid: str, artifact_id: str):
    sess = session_for_request(sid)
    payload = request.get_json(silent=True) or {}
    state = _result_store(sid).patch_artifact(
        artifact_id,
        {"core_selected": bool(payload.get("core_selected", payload.get("selected", True)))},
    )
    if state is None:
        return jsonify({"error": "产物不存在"}), 404
    return jsonify({"ok": True, "state": without_replay_snapshots(state)})


@bp.patch("/api/session/<sid>/research/artifacts/<artifact_id>/paper")
@_exclusive_research_mutation
def patch_research_artifact_paper(sid: str, artifact_id: str):
    sess = session_for_request(sid)
    payload = request.get_json(silent=True) or {}
    updates = {
        "included_in_paper": bool(payload.get("included_in_paper", payload.get("included", True))),
    }
    for field in ("paper_group", "paper_order", "title", "caption", "footnote"):
        if field in payload:
            updates[field] = payload.get(field)
    state = _result_store(sid).patch_artifact(artifact_id, updates)
    if state is None:
        return jsonify({"error": "产物不存在"}), 404
    return jsonify({"ok": True, "state": without_replay_snapshots(state)})


@bp.patch("/api/session/<sid>/research/paper/reorder")
@_exclusive_research_mutation
def reorder_research_paper_items(sid: str):
    sess = session_for_request(sid)
    payload = request.get_json(silent=True) or {}
    items = payload.get("items") if isinstance(payload.get("items"), list) else []
    state = _result_store(sid).reorder_paper(items)
    return jsonify({"ok": True, "state": without_replay_snapshots(state)})


def _bounded_preview_field(value: object) -> tuple[str, bool]:
    text = str(value or "")
    encoded = text.encode("utf-8")
    if len(encoded) <= _BROWSER_PREVIEW_MAX_FIELD_BYTES:
        return text, False
    bounded = encoded[:_BROWSER_PREVIEW_MAX_FIELD_BYTES].decode("utf-8", errors="ignore")
    return bounded, True


def _unique_preview_columns(header: list[str]) -> tuple[list[str], bool]:
    columns: list[str] = []
    counts: dict[str, int] = {}
    truncated = len(header) > _BROWSER_PREVIEW_MAX_COLUMNS
    for index, raw in enumerate(header[:_BROWSER_PREVIEW_MAX_COLUMNS], start=1):
        column, field_truncated = _bounded_preview_field(raw)
        truncated = truncated or field_truncated
        base = column or f"column_{index}"
        count = counts.get(base, 0) + 1
        counts[base] = count
        columns.append(base if count == 1 else f"{base} ({count})")
    return columns, truncated


def _artifact_table_as_json(path: Path):
    try:
        with path.open(newline="", encoding="utf-8-sig") as handle:
            reader = csv.reader(handle, delimiter=_table_delimiter(path))
            header = next(reader, [])
            columns, truncated = _unique_preview_columns(header)
            rows: list[dict[str, str]] = []
            total_rows = 0
            total_columns = len(header)
            for raw_row in reader:
                total_rows += 1
                total_columns = max(total_columns, len(raw_row))
                if len(rows) >= _BROWSER_PREVIEW_MAX_ROWS:
                    truncated = True
                    continue
                values: list[str] = []
                for raw in raw_row[: len(columns)]:
                    value, field_truncated = _bounded_preview_field(raw)
                    truncated = truncated or field_truncated
                    values.append(value)
                values.extend([""] * (len(columns) - len(values)))
                rows.append(dict(zip(columns, values, strict=True)))
    except OSError:
        abort(404)
    except UnicodeDecodeError:
        abort(422, description="表格文件不是 UTF-8 编码，无法预览；请以 UTF-8 编码重新导出。")
    except csv.Error:
        abort(422, description="表格无法安全解析")

    truncated = truncated or total_rows > len(rows) or total_columns > len(columns)
    payload = {
        "ok": True,
        "kind": "table",
        "columns": columns,
        "rows": rows,
        "truncated": truncated,
        "total_rows": total_rows,
        "total_columns": total_columns,
    }
    response = jsonify(payload)
    while rows and len(response.get_data()) > _BROWSER_PREVIEW_MAX_RESPONSE_BYTES:
        rows.pop()
        payload["truncated"] = True
        response = jsonify(payload)
    return response


def _artifact_as_json(path: Path):
    suffix = path.suffix.lower()
    if _is_table_artifact(path):
        return _artifact_table_as_json(path)
    if suffix in {".txt", ".log", ".md"}:
        try:
            with path.open(encoding="utf-8") as handle:
                text = handle.read(200_001)
        except OSError:
            abort(404)
        return jsonify(
            {
                "ok": True,
                "kind": "text",
                "text": text[:200_000],
                "truncated": len(text) > 200_000,
            }
        )
    return jsonify(
        {
            "ok": True,
            "kind": "file",
            "name": path.name,
            "size_bytes": path.stat().st_size,
        }
    )


def _download_name(path: Path, suffix: str) -> str:
    stem = path.stem.strip() or "research-artifact"
    return f"{stem}{suffix}"


def _read_table_rows(path: Path) -> list[list[str]]:
    if not _is_table_artifact(path):
        abort(415)
    max_bytes = _positive_int_env("BAA_RESEARCH_EXPORT_MAX_TABLE_BYTES", 8 * 1024 * 1024)
    max_rows = _positive_int_env("BAA_RESEARCH_EXPORT_MAX_TABLE_ROWS", 10_000)
    max_cells = _positive_int_env("BAA_RESEARCH_EXPORT_MAX_TABLE_CELLS", 100_000)
    try:
        if path.stat().st_size > max_bytes:
            abort(413, description="表格超过导出大小限制")
        with path.open(newline="", encoding="utf-8-sig") as handle:
            rows: list[list[str]] = []
            cell_count = 0
            for row_number, row in enumerate(csv.reader(handle, delimiter=_table_delimiter(path)), start=1):
                if row_number > max_rows:
                    abort(413, description="表格超过导出行数限制")
                cell_count += len(row)
                if cell_count > max_cells:
                    abort(413, description="表格超过导出单元格限制")
                rows.append(row)
            return rows
    except OSError:
        abort(404)


def _safe_spreadsheet_value(value: object) -> object:
    """Neutralize formula-like text while preserving ordinary cell values."""
    if not isinstance(value, str):
        return value
    candidate = value.strip()
    if re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?", candidate):
        return value
    if value.startswith(("=", "+", "-", "@", "\t", "\r")) or value.lstrip(" ").startswith(
        ("=", "+", "-", "@", "\t", "\r")
    ):
        return "'" + value
    return value


def _validate_image_pixel_budget(image) -> None:
    width, height = image.size
    max_pixels = _positive_int_env("BAA_RESEARCH_EXPORT_MAX_IMAGE_PIXELS", 40_000_000)
    if width <= 0 or height <= 0 or width * height > max_pixels:
        raise ValueError("image pixel count exceeds export limit")


def _artifact_table_as_csv(path: Path):
    rows = _read_table_rows(path)
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerows([[_safe_spreadsheet_value(cell) for cell in row] for row in rows])
    payload = io.BytesIO(buffer.getvalue().encode("utf-8"))
    return send_file(
        payload,
        mimetype="text/csv",
        as_attachment=True,
        download_name=_download_name(path, ".csv"),
    )


def _cell_text(row: list[str], index: int) -> str:
    if index >= len(row):
        return ""
    return str(row[index])


def _set_docx_cell_border(cell, **edges):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.find(qn("w:tcBorders"))
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)

    for edge, attrs in edges.items():
        element = tc_borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tc_borders.append(element)
        for key, value in attrs.items():
            element.set(qn(f"w:{key}"), str(value))


def _style_docx_three_line_table(table) -> None:
    no_border = {"val": "nil", "sz": "0", "space": "0", "color": "FFFFFF"}
    rule = {"val": "single", "sz": "8", "space": "0", "color": "111827"}

    for row in table.rows:
        for cell in row.cells:
            _set_docx_cell_border(
                cell,
                top=no_border,
                left=no_border,
                bottom=no_border,
                right=no_border,
                insideH=no_border,
                insideV=no_border,
            )

    if not table.rows:
        return
    for cell in table.rows[0].cells:
        _set_docx_cell_border(cell, top=rule, bottom=rule, left=no_border, right=no_border)
    for cell in table.rows[-1].cells:
        _set_docx_cell_border(cell, bottom=rule, left=no_border, right=no_border)


def _artifact_table_as_docx(path: Path):
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt
    except ImportError:
        abort(501)

    rows = _read_table_rows(path)
    if not rows:
        rows = [["No data"]]
    column_count = max(len(row) for row in rows) or 1

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run(f"表格：{path.stem.replace('_', ' ')}")
    run.bold = True
    run.font.size = Pt(11)

    table = document.add_table(rows=0, cols=column_count)
    table.autofit = True
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for column_index in range(column_count):
            cell = cells[column_index]
            cell.text = _cell_text(row, column_index)
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for paragraph_run in paragraph.runs:
                    paragraph_run.font.size = Pt(9)
                    paragraph_run.bold = row_index == 0
    _style_docx_three_line_table(table)

    note = document.add_paragraph()
    note_run = note.add_run("注：本表由已执行的 R 产物导出，数值保持与源表一致。")
    note_run.italic = True
    note_run.font.size = Pt(8)

    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return send_file(
        buffer,
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        as_attachment=True,
        download_name=_download_name(path, ".docx"),
    )


def _artifact_table_as_xlsx(path: Path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        abort(501)

    rows = _read_table_rows(path)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "表格"

    for row in rows:
        worksheet.append([_safe_spreadsheet_value(cell) for cell in row])

    if rows:
        rule = Side(style="thin", color="111827")
        empty = Side(style=None)
        header_fill = PatternFill("solid", fgColor="F8FAFC")
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.border = Border(top=rule, bottom=rule, left=empty, right=empty)
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        for cell in worksheet[worksheet.max_row]:
            cell.border = Border(bottom=rule, left=empty, right=empty)
        for row in worksheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for column_cells in worksheet.columns:
            column_index = column_cells[0].column
            width = max(len(str(cell.value or "")) for cell in column_cells) + 2
            worksheet.column_dimensions[get_column_letter(column_index)].width = max(10, min(width, 48))

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return send_file(
        buffer,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=_download_name(path, ".xlsx"),
    )


def _companion_artifact(path: Path, suffix: str) -> Path | None:
    candidate = path.with_suffix(suffix)
    try:
        resolved = candidate.resolve(strict=True)
        resolved.relative_to(path.parent.resolve(strict=False))
    except (OSError, RuntimeError, ValueError):
        return None
    return resolved if resolved.is_file() else None


def _artifact_image_as_pdf(path: Path):
    native_pdf = _companion_artifact(path, ".pdf")
    if native_pdf is not None:
        return send_file(
            native_pdf,
            mimetype="application/pdf",
            as_attachment=True,
            download_name=_download_name(path, ".pdf"),
        )

    if path.suffix.lower() == ".pdf":
        return send_file(
            path,
            mimetype="application/pdf",
            as_attachment=True,
            download_name=_download_name(path, ".pdf"),
        )
    if path.suffix.lower() not in {".png", ".jpg", ".jpeg"}:
        abort(415)
    if path.stat().st_size > _positive_int_env("BAA_RESEARCH_EXPORT_MAX_FIGURE_BYTES", 64 * 1024 * 1024):
        abort(413, description="图形超过转换大小限制")
    try:
        from PIL import Image

        with Image.open(path) as image:
            try:
                _validate_image_pixel_budget(image)
            except ValueError:
                abort(413, description="图形像素数量超过转换限制")
            rgb = image.convert("RGB")
            buffer = io.BytesIO()
            rgb.save(buffer, format="PDF", resolution=300.0)
    except OSError:
        abort(404)
    except ImportError:
        abort(501)
    buffer.seek(0)
    return send_file(
        buffer,
        mimetype="application/pdf",
        as_attachment=True,
        download_name=_download_name(path, ".pdf"),
    )


def _artifact_image_as_tiff(path: Path):
    native_tiff = _companion_artifact(path, ".tiff") or _companion_artifact(path, ".tif")
    if native_tiff is not None:
        return send_file(
            native_tiff,
            mimetype="image/tiff",
            as_attachment=True,
            download_name=_download_name(path, native_tiff.suffix),
        )

    if path.suffix.lower() in {".tif", ".tiff"}:
        return send_file(
            path,
            mimetype="image/tiff",
            as_attachment=True,
            download_name=_download_name(path, path.suffix),
        )
    if path.suffix.lower() not in {".png", ".jpg", ".jpeg"}:
        abort(415)
    if path.stat().st_size > _positive_int_env("BAA_RESEARCH_EXPORT_MAX_FIGURE_BYTES", 64 * 1024 * 1024):
        abort(413, description="图形超过转换大小限制")
    try:
        from PIL import Image

        with Image.open(path) as image:
            try:
                _validate_image_pixel_budget(image)
            except ValueError:
                abort(413, description="图形像素数量超过转换限制")
            rgb = image.convert("RGB")
            buffer = io.BytesIO()
            rgb.save(buffer, format="TIFF", compression="tiff_lzw", dpi=(600, 600))
    except OSError:
        abort(404)
    except ImportError:
        abort(501)
    buffer.seek(0)
    return send_file(
        buffer,
        mimetype="image/tiff",
        as_attachment=True,
        download_name=_download_name(path, ".tiff"),
    )


def _safe_export_name(value: str, fallback: str = "research_export") -> str:
    name = secure_filename(str(value or "").strip())
    return name or fallback


def _load_v2_state_or_404(sid: str, sess) -> dict:
    state, exists, _saved_at = _result_store(sid).load()
    if not exists:
        abort(404)
    return state


def _state_artifacts(state: dict, *, included_only: bool = False, kind: str = "") -> list[dict]:
    artifacts: list[dict] = []
    for turn in state.get("analysis_turns", []):
        for artifact in turn.get("artifacts", []):
            if included_only and not artifact.get("included_in_paper"):
                continue
            if kind and artifact.get("kind") != kind:
                continue
            artifacts.append({**artifact, "_turn": turn})
    artifacts.sort(
        key=lambda item: (
            str(item.get("paper_group") or ""),
            int(item.get("paper_order") or 0),
            int(item.get("priority") or 999),
            str(item.get("title") or ""),
        )
    )
    return artifacts


def _artifact_path_or_none(artifact: dict, allowed_roots: list[Path]) -> Path | None:
    raw = str(artifact.get("path") or "")
    if not raw or "\x00" in raw:
        return None
    try:
        path = Path(raw).resolve(strict=True)
    except (OSError, RuntimeError):
        return None
    if not path.is_file():
        return None
    for root in allowed_roots:
        try:
            path.relative_to(root.resolve(strict=False))
            return path
        except ValueError:
            continue
    return None


def _validate_export_budget(
    state: dict,
    allowed_roots: list[Path],
    *,
    validate_tables: bool = True,
    validate_figures: bool = True,
) -> None:
    artifacts = _state_artifacts(state, included_only=True)
    max_artifacts = _positive_int_env("BAA_RESEARCH_EXPORT_MAX_ARTIFACTS", 200)
    max_source_bytes = _positive_int_env("BAA_RESEARCH_EXPORT_MAX_SOURCE_BYTES", 256 * 1024 * 1024)
    max_total_rows = _positive_int_env("BAA_RESEARCH_EXPORT_MAX_TOTAL_ROWS", 10_000)
    max_total_cells = _positive_int_env("BAA_RESEARCH_EXPORT_MAX_TOTAL_CELLS", 100_000)
    max_figure_pixels = _positive_int_env("BAA_RESEARCH_EXPORT_MAX_IMAGE_PIXELS", 40_000_000)
    max_total_figure_pixels = _positive_int_env("BAA_RESEARCH_EXPORT_MAX_TOTAL_FIGURE_PIXELS", 80_000_000)
    if len(artifacts) > max_artifacts:
        abort(413, description="纳入导出的科研产物过多")
    total_bytes = 0
    total_rows = 0
    total_cells = 0
    total_figure_pixels = 0
    table_shapes: dict[Path, tuple[int, int]] = {}
    figure_shapes: dict[Path, int] = {}
    for artifact in artifacts:
        path = _artifact_path_or_none(artifact, allowed_roots)
        if path is None:
            continue
        try:
            # Repeated references are materialized repeatedly by the builders.
            total_bytes += path.stat().st_size
        except OSError:
            continue
        if total_bytes > max_source_bytes:
            abort(413, description="导出源文件总量超过限制")
        if validate_tables and _is_table_artifact(path):
            if path not in table_shapes:
                rows = _read_table_rows(path)
                table_shapes[path] = (len(rows), sum(len(row) for row in rows))
            row_count, cell_count = table_shapes[path]
            # Builders expand every artifact entry, even when several entries
            # reference the same source file, so the budget must do the same.
            total_rows += row_count
            total_cells += cell_count
            if total_rows > max_total_rows:
                abort(413, description="导出表格累计行数超过限制")
            if total_cells > max_total_cells:
                abort(413, description="导出表格累计单元格数超过限制")
        if validate_figures and path.suffix.lower() in {
            ".png",
            ".jpg",
            ".jpeg",
            ".tif",
            ".tiff",
            ".bmp",
            ".webp",
        }:
            if path not in figure_shapes:
                try:
                    from PIL import Image

                    with Image.open(path) as image:
                        figure_shapes[path] = int(image.width) * int(image.height)
                except OSError:
                    figure_shapes[path] = 0
            figure_pixels = figure_shapes[path]
            if figure_pixels > max_figure_pixels:
                abort(413, description="单个导出图形解码像素超过限制")
            total_figure_pixels += figure_pixels
            if total_figure_pixels > max_total_figure_pixels:
                abort(413, description="导出图形累计解码像素超过限制")


def _docx_add_table(document, artifact: dict, allowed_roots: list[Path]) -> None:
    from docx.shared import Pt

    path = _artifact_path_or_none(artifact, allowed_roots)
    if path is None or not _is_table_artifact(path):
        return
    rows = _read_table_rows(path)
    if not rows:
        return
    title = artifact.get("title") or path.stem.replace("_", " ")
    document.add_paragraph(str(title), style=None).runs[0].bold = True
    column_count = max(len(row) for row in rows) or 1
    table = document.add_table(rows=0, cols=column_count)
    table.autofit = True
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for col_index in range(column_count):
            cell = cells[col_index]
            cell.text = _cell_text(row, col_index)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    run.bold = row_index == 0
    _style_docx_three_line_table(table)
    footnote = str(artifact.get("footnote") or "").strip()
    if footnote:
        note = document.add_paragraph(f"注：{footnote}")
        for run in note.runs:
            run.font.size = Pt(8)
            run.italic = True


def _artifact_text_for_export(
    artifact: dict,
    allowed_roots: list[Path],
    *,
    max_bytes: int | None = None,
) -> tuple[str, bool]:
    """Read a bounded, XML-safe textual artifact for Word/package export."""

    limit = max_bytes or _positive_int_env("BAA_RESEARCH_EXPORT_MAX_TEXT_BYTES", 1 * 1024 * 1024)
    path = _artifact_path_or_none(artifact, allowed_roots)
    text = ""
    truncated = False
    if path is not None and path.suffix.lower() in {".txt", ".log", ".md", ".json", ".yml", ".yaml"}:
        try:
            with path.open(encoding="utf-8", errors="replace") as handle:
                text = handle.read(limit + 1)
            truncated = len(text.encode("utf-8")) > limit
            if truncated:
                encoded = text.encode("utf-8")[:limit]
                text = encoded.decode("utf-8", errors="ignore")
        except OSError:
            text = ""
    if not text:
        text = str(artifact.get("preview") or "")
    xml_safe = "".join(character for character in text if character in "\t\n\r" or ord(character) >= 32)
    return xml_safe.strip(), truncated


def _build_word_export(state: dict, allowed_roots: list[Path]) -> io.BytesIO:
    try:
        from docx import Document
        from docx.shared import Inches, Pt
    except ImportError:
        abort(501)

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)

    dataset = state.get("dataset") or {}
    document.add_heading("科研统计结果草稿", level=1)
    if dataset.get("name"):
        document.add_paragraph(f"数据集：{dataset.get('name')}")
    document.add_paragraph(f"生成时间：{datetime.now().isoformat(timespec='seconds')}")

    document.add_heading("结果解释", level=2)
    narratives = [
        str(turn.get("narrative") or "").strip()
        for turn in state.get("analysis_turns", [])
        if str(turn.get("narrative") or "").strip()
    ]
    if narratives:
        for narrative in narratives:
            document.add_paragraph(narrative)
    else:
        document.add_paragraph("尚未纳入可用的论文结果解释。")

    tables = _state_artifacts(state, included_only=True, kind="table")
    document.add_heading("表格", level=2)
    if tables:
        for artifact in tables:
            _docx_add_table(document, artifact, allowed_roots)
            document.add_paragraph("")
    else:
        document.add_paragraph("尚未纳入论文表格。")

    figures = _state_artifacts(state, included_only=True, kind="figure")
    document.add_heading("图形", level=2)
    if figures:
        for artifact in figures:
            title = artifact.get("title") or artifact.get("name") or "Figure"
            document.add_paragraph(str(title)).runs[0].bold = True
            path = _artifact_path_or_none(artifact, allowed_roots)
            if path is not None and path.suffix.lower() in {".png", ".jpg", ".jpeg"}:
                try:
                    document.add_picture(str(path), width=Inches(5.8))
                except Exception:
                    document.add_paragraph(f"图形文件：{path.name}")
            caption = str(artifact.get("caption") or "").strip()
            if caption:
                document.add_paragraph(f"图注：{caption}")
    else:
        document.add_paragraph("尚未纳入论文图形。")

    models = _state_artifacts(state, included_only=True, kind="model")
    document.add_heading("模型结果", level=2)
    if models:
        for artifact in models:
            title = artifact.get("title") or artifact.get("name") or "Model"
            document.add_paragraph(str(title)).runs[0].bold = True
            content, truncated = _artifact_text_for_export(artifact, allowed_roots)
            if content:
                document.add_paragraph(content)
                if truncated:
                    document.add_paragraph("模型摘要过长，Word 中仅纳入前部内容；完整文件见交付包。")
            else:
                path = _artifact_path_or_none(artifact, allowed_roots)
                document.add_paragraph(
                    f"模型文件：{path.name}" if path is not None else "模型文件当前不可读取。"
                )
            caption = str(artifact.get("caption") or "").strip()
            if caption:
                document.add_paragraph(f"说明：{caption}")
    else:
        document.add_paragraph("尚未纳入模型结果。")

    method_notes = [
        str(turn.get("methods_note") or "").strip()
        for turn in state.get("analysis_turns", [])
        if str(turn.get("methods_note") or "").strip()
    ]
    document.add_heading("方法说明", level=2)
    if method_notes:
        for note in method_notes:
            document.add_paragraph(note)
    else:
        document.add_paragraph("本导出包未包含额外方法说明。")

    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer


def _build_excel_export(state: dict, allowed_roots: list[Path]) -> io.BytesIO:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        abort(501)

    workbook = Workbook()
    metadata = workbook.active
    metadata.title = "元数据"
    metadata.append(["字段", "值"])
    metadata.append(["数据集", _safe_spreadsheet_value((state.get("dataset") or {}).get("name", ""))])
    metadata.append(["生成时间", datetime.now().isoformat(timespec="seconds")])
    metadata.append(["分析轮次", len(state.get("analysis_turns", []))])

    tables = _state_artifacts(state, included_only=True, kind="table")
    if not tables:
        worksheet = workbook.create_sheet("无纳入表格")
        worksheet.append(["当前没有纳入论文的表格"])

    for index, artifact in enumerate(tables, start=1):
        path = _artifact_path_or_none(artifact, allowed_roots)
        if path is None or not _is_table_artifact(path):
            continue
        title = str(artifact.get("title") or f"表格 {index}")[:26]
        sheet_name = _safe_export_name(title, f"table_{index}")[:26] or f"table_{index}"
        worksheet = workbook.create_sheet(f"{index}_{sheet_name}"[:31])
        rows = _read_table_rows(path)
        for row in rows:
            worksheet.append([_safe_spreadsheet_value(cell) for cell in row])
        if rows:
            fill = PatternFill("solid", fgColor="F8FAFC")
            for cell in worksheet[1]:
                cell.font = Font(bold=True)
                cell.fill = fill
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            for row in worksheet.iter_rows():
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
            for column_cells in worksheet.columns:
                column_index = column_cells[0].column
                width = max(len(str(cell.value or "")) for cell in column_cells) + 2
                worksheet.column_dimensions[get_column_letter(column_index)].width = max(10, min(width, 50))
        if artifact.get("footnote"):
            worksheet.append([])
            worksheet.append(["脚注", _safe_spreadsheet_value(artifact.get("footnote"))])

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def _figure_bytes_for_format(path: Path, fmt: str) -> tuple[bytes, str]:
    max_bytes = _positive_int_env("BAA_RESEARCH_EXPORT_MAX_FIGURE_BYTES", 64 * 1024 * 1024)
    if path.stat().st_size > max_bytes:
        raise ValueError("figure exceeds export size limit")
    fmt = fmt.lower()
    if fmt == "png" and path.suffix.lower() == ".png":
        return path.read_bytes(), ".png"
    if fmt == "pdf":
        companion = _companion_artifact(path, ".pdf")
        if companion is not None:
            if companion.stat().st_size > max_bytes:
                raise ValueError("figure companion exceeds export size limit")
            return companion.read_bytes(), ".pdf"
        from PIL import Image

        with Image.open(path) as image:
            _validate_image_pixel_budget(image)
            rgb = image.convert("RGB")
            buffer = io.BytesIO()
            rgb.save(buffer, format="PDF", resolution=300.0)
            return buffer.getvalue(), ".pdf"
    if fmt == "tiff":
        companion = _companion_artifact(path, ".tiff") or _companion_artifact(path, ".tif")
        if companion is not None:
            if companion.stat().st_size > max_bytes:
                raise ValueError("figure companion exceeds export size limit")
            return companion.read_bytes(), companion.suffix
        from PIL import Image

        with Image.open(path) as image:
            _validate_image_pixel_budget(image)
            rgb = image.convert("RGB")
            buffer = io.BytesIO()
            rgb.save(buffer, format="TIFF", compression="tiff_lzw", dpi=(600, 600))
            return buffer.getvalue(), ".tiff"
    if fmt == "png":
        from PIL import Image

        with Image.open(path) as image:
            _validate_image_pixel_budget(image)
            buffer = io.BytesIO()
            image.save(buffer, format="PNG")
            return buffer.getvalue(), ".png"
    raise ValueError(f"unsupported figure format: {fmt}")


def _build_figures_zip(state: dict, allowed_roots: list[Path]) -> io.BytesIO:
    figures = _state_artifacts(state, included_only=True, kind="figure")
    manifest: list[dict] = []
    buffer = io.BytesIO()
    max_output_bytes = _positive_int_env(
        "BAA_RESEARCH_EXPORT_MAX_TOTAL_FIGURE_OUTPUT_BYTES", 256 * 1024 * 1024
    )
    total_output_bytes = 0
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for index, artifact in enumerate(figures, start=1):
            path = _artifact_path_or_none(artifact, allowed_roots)
            if path is None:
                continue
            stem = _safe_export_name(
                artifact.get("title") or artifact.get("name") or f"figure_{index}", f"figure_{index}"
            )
            item_manifest = {
                "artifact_id": artifact.get("artifact_id"),
                "title": artifact.get("title"),
                "caption": artifact.get("caption"),
                "files": [],
            }
            for fmt in ("png", "pdf", "tiff"):
                try:
                    payload, suffix = _figure_bytes_for_format(path, fmt)
                except Exception:
                    continue
                total_output_bytes += len(payload)
                if total_output_bytes > max_output_bytes:
                    abort(413, description="导出图形生成后的累计文件大小超过限制")
                filename = f"figures/{index:02d}_{stem}{suffix}"
                archive.writestr(filename, payload)
                item_manifest["files"].append(filename)
            manifest.append(item_manifest)
        archive.writestr("figures_manifest.json", json.dumps(manifest, ensure_ascii=False, indent=2))
    buffer.seek(0)
    return buffer


def _build_models_zip(state: dict, allowed_roots: list[Path]) -> io.BytesIO:
    models = _state_artifacts(state, included_only=True, kind="model")
    manifest: list[dict] = []
    buffer = io.BytesIO()
    max_output_bytes = _positive_int_env(
        "BAA_RESEARCH_EXPORT_MAX_TOTAL_MODEL_BYTES", 64 * 1024 * 1024
    )
    total_output_bytes = 0
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for index, artifact in enumerate(models, start=1):
            path = _artifact_path_or_none(artifact, allowed_roots)
            if path is None:
                continue
            try:
                payload = path.read_bytes()
            except OSError:
                continue
            total_output_bytes += len(payload)
            if total_output_bytes > max_output_bytes:
                abort(413, description="导出模型文件累计大小超过限制")
            stem = _safe_export_name(
                artifact.get("title") or artifact.get("name") or f"model_{index}",
                f"model_{index}",
            )
            suffix = path.suffix if path.suffix else ".bin"
            filename = f"models/{index:02d}_{stem}{suffix}"
            archive.writestr(filename, payload)
            manifest.append(
                {
                    "artifact_id": artifact.get("artifact_id"),
                    "title": artifact.get("title"),
                    "caption": artifact.get("caption"),
                    "file": filename,
                }
            )
        archive.writestr("models_manifest.json", json.dumps(manifest, ensure_ascii=False, indent=2))
    buffer.seek(0)
    return buffer


def _delivery_markdown(state: dict) -> str:
    delivery = state.get("delivery_status") or {}
    lines = [
        "# 交付验收清单",
        "",
        f"- 数据集：{(state.get('dataset') or {}).get('name') or '未上传'}",
        f"- 分析轮次：{(state.get('overview') or {}).get('analysis_turn_count') or 0}",
        f"- 表格：{(state.get('overview') or {}).get('table_count') or 0}",
        f"- 图形：{(state.get('overview') or {}).get('figure_count') or 0}",
        f"- 模型：{(state.get('overview') or {}).get('model_count') or 0}",
        f"- 纳入论文产物：{(state.get('overview') or {}).get('included_in_paper_count') or 0}",
        "",
        "## 状态",
    ]
    for item in delivery.get("items") or []:
        lines.append(
            f"- {'通过' if item.get('ok') else '待处理'}：{item.get('label')}（{item.get('detail') or ''}）"
        )
    return "\n".join(lines) + "\n"


def _public_delivery_manifest(state):
    """Strip server filesystem capabilities from the portable manifest."""
    return public_projection(
        state,
        drop_keys={
            "path",
            "stored_path",
            "artifacts_dir",
            "file_path",
            "source_path",
            "output_path",
            "workdir",
            "cwd",
            "stdout",
            "stderr",
            "code",
            "execution_snapshot",
        },
    )


def _build_delivery_package(state: dict, allowed_roots: list[Path]) -> io.BytesIO:
    from data.reproducibility import write_reproducibility_bundle

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(
            "analysis_manifest.json",
            json.dumps(_public_delivery_manifest(state), ensure_ascii=False, indent=2),
        )
        archive.writestr("delivery_checklist.md", _delivery_markdown(state))
        archive.writestr("research_results.docx", _build_word_export(state, allowed_roots).getvalue())
        archive.writestr("research_tables.xlsx", _build_excel_export(state, allowed_roots).getvalue())
        archive.writestr("research_figures.zip", _build_figures_zip(state, allowed_roots).getvalue())
        archive.writestr("research_models.zip", _build_models_zip(state, allowed_roots).getvalue())
        write_reproducibility_bundle(archive, state)
    buffer.seek(0)
    return buffer


@bp.post("/api/session/<sid>/research/export/word")
def export_research_word(sid: str):
    sess = session_for_request(sid)
    state = _load_v2_state_or_404(sid, sess)
    allowed_roots = _allowed_artifact_roots(sid, sess)
    _validate_export_budget(state, allowed_roots)
    buffer = _build_word_export(state, allowed_roots)
    return send_file(
        buffer,
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        as_attachment=True,
        download_name="research_results.docx",
    )


@bp.post("/api/session/<sid>/research/export/excel")
def export_research_excel(sid: str):
    sess = session_for_request(sid)
    state = _load_v2_state_or_404(sid, sess)
    allowed_roots = _allowed_artifact_roots(sid, sess)
    _validate_export_budget(state, allowed_roots, validate_figures=False)
    buffer = _build_excel_export(state, allowed_roots)
    return send_file(
        buffer,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="research_tables.xlsx",
    )


@bp.post("/api/session/<sid>/research/export/figures")
def export_research_figures(sid: str):
    sess = session_for_request(sid)
    state = _load_v2_state_or_404(sid, sess)
    allowed_roots = _allowed_artifact_roots(sid, sess)
    _validate_export_budget(state, allowed_roots, validate_tables=False)
    buffer = _build_figures_zip(state, allowed_roots)
    return send_file(
        buffer,
        mimetype="application/zip",
        as_attachment=True,
        download_name="research_figures.zip",
    )


@bp.post("/api/session/<sid>/research/export/package")
def export_research_package(sid: str):
    sess = session_for_request(sid)
    state = _load_v2_state_or_404(sid, sess)
    allowed_roots = _allowed_artifact_roots(sid, sess)
    _validate_export_budget(state, allowed_roots)
    try:
        buffer = _build_delivery_package(state, allowed_roots)
    except OverflowError as exc:
        abort(413, description=str(exc))
    return send_file(
        buffer,
        mimetype="application/zip",
        as_attachment=True,
        download_name="research_delivery_package.zip",
    )


@bp.post("/api/session/<sid>/research/upload")
@_exclusive_research_mutation
def upload_research_dataset(sid: str):
    sess = session_for_request(sid)
    files = request.files.getlist("file")
    valid_files = [upload for upload in files if upload and upload.filename]
    if len(valid_files) != 1 or len(files) != 1:
        return jsonify({"error": "每次必须且只能上传一个数据文件"}), 400

    upload = valid_files[0]
    max_file_bytes = _research_upload_limit_bytes()
    suffix = Path(upload.filename).suffix
    safe_name = secure_filename(upload.filename) or ""
    # secure_filename strips every non-ASCII char, so a Chinese-named file
    # ("临床数据.csv") degrades to "csv" with no extension at all.
    if suffix and not safe_name.endswith(suffix):
        safe_name = f"{safe_name or 'upload'}{suffix}"
    safe_name = safe_name or f"upload{suffix}"
    workspace = default_research_workspace()
    try:
        with tempfile.TemporaryDirectory(
            prefix="upload-",
            dir=workspace.browser_upload_staging_root(),
        ) as tmp:
            tmp_path = Path(tmp) / safe_name
            expected_bytes = int(request.content_length or upload.content_length or 0)
            upload_bytes = _stream_browser_upload(
                upload,
                tmp_path,
                expected_bytes=expected_bytes,
                limit=max_file_bytes,
            )
            if upload_bytes <= 0:
                raise ValueError("上传的数据文件为空或无法读取")
            dataset = workspace.replace_staged_upload(sid, tmp_path, upload.filename)
    except OSError as exc:
        if exc.errno in {errno.ENOSPC, getattr(errno, "EDQUOT", errno.ENOSPC)}:
            detail = "本地数据存储空间不足，上传未写入"
        else:
            detail = "无法创建本地上传暂存文件"
        return jsonify({"error": f"{upload.filename}: {detail}"}), 400
    except (RuntimeError, ValueError) as exc:
        return jsonify({"error": f"{upload.filename}: {exc}"}), 400

    public_dataset = dataset.to_dict()
    sess.research_dataset_id = dataset.dataset_id

    return jsonify(
        {
            "ok": True,
            "dataset": public_dataset,
            "datasets": [public_dataset],
            "errors": [],
        }
    )


@bp.post("/api/session/<sid>/research/import-local")
@_exclusive_research_mutation
def import_local_research_dataset(sid: str):
    """Copy native-selected files/directories into app-owned storage.

    Arbitrary paths are never accepted from HTTP.  The Qt bridge issues a
    short-lived, one-use capability for the exact selection and only the
    loopback desktop runtime enables this endpoint. Native selection and
    browser drag/drop intentionally share the same durable copy semantics.
    """
    if os.environ.get("BAA_DESKTOP_MODE") != "1":
        return jsonify({"error": "本地科研数据导入仅在桌面应用中可用"}), 403

    payload = request.get_json(silent=True) or {}
    token = str(payload.get("grant_token") or "").strip()
    try:
        from desktop.local_import import default_local_import_grants

        grant = default_local_import_grants().consume(token)
        assets = infer_scientific_asset_specs(grant.paths)
        dataset = default_research_workspace().replace_assets(
            sid,
            assets,
            access_mode="copy",
        )
    except (RuntimeError, ValueError) as exc:
        return jsonify({"error": str(exc)}), 400

    public_dataset = dataset.to_dict()
    sess = session_for_request(sid)
    sess.research_dataset_id = dataset.dataset_id
    return jsonify(
        {
            "ok": True,
            "dataset": public_dataset,
            "datasets": [public_dataset],
            "errors": [],
            "import_mode": "local_copy",
        }
    )


@bp.get("/api/session/<sid>/research/storage")
def research_workspace_storage(sid: str):
    session_for_request(sid)
    return jsonify({"ok": True, "storage": default_research_workspace().storage_usage(sid)})


@bp.delete("/api/session/<sid>/research/workspace")
@_exclusive_research_mutation
def delete_research_workspace(sid: str):
    sess = session_for_request(sid)
    workspace = default_research_workspace()
    removed = workspace.delete_current(sid)
    sess.research_dataset_id = ""
    sess.recent_sql = []
    sess.recent_artifacts = []
    sess.turn_activations = []
    return jsonify(
        {
            "ok": True,
            "removed": bool(removed),
            "storage": workspace.storage_usage(sid),
        }
    )


@bp.get("/api/session/<sid>/research/datasets")
def list_research_datasets(sid: str):
    sess = session_for_request(sid)
    dataset = default_research_workspace().current(sid)
    datasets = [dataset] if dataset is not None else []
    active_dataset_id = dataset.dataset_id if dataset is not None else ""
    sess.research_dataset_id = active_dataset_id
    return jsonify(
        {
            "ok": True,
            "datasets": [dataset.to_dict() for dataset in datasets],
            "active_dataset_id": active_dataset_id,
        }
    )


@bp.patch("/api/session/<sid>/research/datasets/active")
@_exclusive_research_mutation
def set_active_research_dataset(sid: str):
    sess = session_for_request(sid)
    payload = request.get_json(silent=True) or {}
    dataset_id = str(payload.get("dataset_id") or "").strip()
    dataset = default_research_workspace().current(sid)
    if dataset is None:
        return jsonify({"error": "数据集不存在"}), 404
    if dataset_id != dataset.dataset_id:
        return jsonify({"error": "当前版本仅允许一个活动数据集，请上传新文件进行替换。"}), 409
    sess.research_dataset_id = dataset.dataset_id
    state, _exists, _saved_at = _result_store(sid).load()
    return jsonify(
        {
            "ok": True,
            "active_dataset_id": dataset.dataset_id,
            "dataset": dataset.to_dict(),
            "state": without_replay_snapshots(state),
        }
    )


@bp.get("/api/session/<sid>/research/datasets/<dataset_id>")
def get_research_dataset(sid: str, dataset_id: str):
    sess = session_for_request(sid)
    dataset = default_research_workspace().current(sid)
    if dataset is None or dataset.dataset_id != dataset_id:
        return jsonify({"error": "数据集不存在"}), 404
    return jsonify({"ok": True, "dataset": dataset.to_dict()})


@bp.post("/api/session/<sid>/research/datasets/<dataset_id>/inspect")
@_exclusive_research_mutation
def inspect_research_dataset(sid: str, dataset_id: str):
    sess = session_for_request(sid)
    store = _store()
    workspace = default_research_workspace()
    dataset = workspace.current(sid)
    if dataset is None:
        return jsonify({"error": "数据集不存在"}), 404
    if dataset.dataset_id != dataset_id:
        return jsonify({"error": "只能理解当前活动数据集，请上传新文件进行替换。"}), 409
    current_state, _state_exists, _saved_at = _result_store(sid).load()
    active_dataset_id = str((current_state.get("dataset") or {}).get("dataset_id") or "") or str(
        getattr(sess, "research_dataset_id", "") or ""
    )
    if active_dataset_id and active_dataset_id != dataset.dataset_id:
        return jsonify({"error": "请先通过数据集选择器切换数据集，再执行数据理解。"}), 409
    sess.research_dataset_id = dataset.dataset_id

    dataset_ref = _dataset_ref(dataset)
    backend = get_r_backend()
    try:
        dataset_path = Path(dataset_ref.path)
        resource_policy = compute_local_resource_policy(
            environment=getattr(backend, "runtime_env", {}),
            storage_path=dataset_path.parent,
            workload_size_bytes=dataset_path.stat().st_size,
        )
        lazy_dataset = dataset_path.stat().st_size > resource_policy.in_memory_threshold_bytes
    except (OSError, TypeError, ValueError):
        lazy_dataset = False
    inspection_packages = dataset_loader_packages(
        dataset_ref.format,
        (asset.format for asset in dataset_ref.assets),
        inspection=True,
        lazy=lazy_dataset,
    )
    dependency_gate = inspect_dependency_gate(
        packages=inspection_packages,
        intent="数据理解与格式读取",
        mode="exact",
        session_id=sid,
        backend=backend,
    )
    if dependency_gate["status"] != "ready":
        return jsonify(
            {
                "ok": False,
                "error": str(dependency_gate.get("headline") or "数据读取依赖尚未就绪"),
                "code": "dependency_preflight",
                "dependency_gate": dependency_gate,
                "dataset": dataset.to_dict(),
            }
        ), 409

    inspection_job_id = ""
    try:
        inspection_job_id = sess.job_runner.begin_tracked(
            "dataset_inspection",
            label="dataset_inspection",
        )
        research_run_controller.begin(inspection_job_id)
        inspection_request = dataset_inspection_request(
            dataset_ref,
            session_id=sid,
            workspace_generation=str(
                (workspace.current_record(sid) or {}).get("workspace_generation") or ""
            ),
            conversation_job_id=inspection_job_id,
        )
        result = run_r_request(backend, inspection_request)
        canceled = (
            research_run_controller.is_cancelled(inspection_job_id)
            or sess.cancel_requested_for(inspection_job_id)
        )
    except JobLimitExceeded as exc:
        return jsonify({"error": str(exc), "code": "job_limit_exceeded"}), 429
    except Exception:
        if inspection_job_id:
            sess.job_runner.fail_tracked(inspection_job_id, "Dataset inspection failed.")
        raise
    finally:
        if inspection_job_id:
            research_run_controller.finish(inspection_job_id)
            sess.clear_cancel(inspection_job_id)

    if canceled:
        sess.job_runner.cancel_tracked(inspection_job_id)
        return jsonify(
            {
                "ok": False,
                "error": "数据理解已停止。",
                "code": "analysis_canceled",
                "job_id": inspection_job_id,
            }
        ), 409
    summary = inspection_summary(result) if result.ok else {}
    public_summary = public_projection(
        summary,
        drop_keys={"local_semantics", "analysis_context", "role_memory"},
    )
    if result.ok:
        analysis_context = (
            dataset.profile.get("analysis_context")
            if isinstance(dataset.profile, dict) and isinstance(dataset.profile.get("analysis_context"), dict)
            else {}
        )
        dataset.profile = {
            "inspection": result.to_dict(),
            "summary": summary,
            "analysis_context": analysis_context,
        }
        store.save(dataset)
        _result_store(sid).set_dataset(dataset.to_dict(), public_summary)
        sess.job_runner.succeed_tracked(
            inspection_job_id,
            {"dataset_id": dataset.dataset_id, "rows": summary.get("rows", 0)},
        )

    status = 200 if result.ok else 500
    public_error = ""
    if not result.ok:
        error_id = uuid.uuid4().hex[:12]
        public_error = f"数据理解失败（错误编号：{error_id}）。"
        log.error(
            "[research] dataset inspection failed error_id=%s dataset_id=%s error=%s stderr=%s",
            error_id,
            dataset.dataset_id,
            result.error,
            result.stderr,
        )
        sess.job_runner.fail_tracked(inspection_job_id, "Dataset inspection failed.")
    public_result = public_r_result(result.to_dict(), error_message=public_error)
    return jsonify(
        {
            "ok": result.ok,
            "error": public_error,
            "dataset": dataset.to_dict(),
            "result": public_result,
            "summary": public_summary,
            "job_id": inspection_job_id,
        }
    ), status


@bp.get("/api/session/<sid>/research/artifact")
def read_research_artifact(sid: str):
    sess = session_for_request(sid)
    path = _resolve_artifact_path(sid, request.args.get("path", ""), sess)
    if path is None:
        abort(404)
    requested_format = (request.args.get("format") or "").strip().lower()
    if requested_format == "json":
        return _artifact_as_json(path)
    if requested_format in {"word", "doc", "docx"}:
        return _artifact_table_as_docx(path)
    if requested_format == "xlsx":
        return _artifact_table_as_xlsx(path)
    if requested_format == "pdf":
        return _artifact_image_as_pdf(path)
    if requested_format in {"tif", "tiff"}:
        return _artifact_image_as_tiff(path)
    if requested_format == "csv" and _is_table_artifact(path):
        return _artifact_table_as_csv(path)

    mimetype = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
    download_name = path.name
    if requested_format == "png" and path.suffix.lower() == ".png":
        mimetype = "image/png"
        download_name = _download_name(path, ".png")
    elif requested_format == "svg" and path.suffix.lower() == ".svg":
        mimetype = "image/svg+xml"
        download_name = _download_name(path, ".svg")
    elif requested_format in {"tif", "tiff"} and path.suffix.lower() in {".tif", ".tiff"}:
        mimetype = "image/tiff"
        download_name = _download_name(path, path.suffix)
    force_attachment = path.suffix.lower() in {".html", ".htm", ".svg", ".xml"}
    if force_attachment and not requested_format:
        mimetype = "application/octet-stream"
    return send_file(
        path,
        mimetype=mimetype,
        as_attachment=force_attachment or request.args.get("download") == "1" or bool(requested_format),
        download_name=download_name,
    )
